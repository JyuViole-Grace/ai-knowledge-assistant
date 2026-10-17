import logging
import re
from pathlib import Path
from .config import SUPPORTED_EXTENSIONS, stable_id, log
from .schemas import Document, Chunk


# A word split across lines by hyphenation: "prin-\ncipe" or "prin -\ncipe"
# (the PDF used here puts a space before the hyphen) -> "principe".
# Only a hyphen that ends a line between two letters is joined, so mid-line
# dashes ("word - word") and digit ranges ("1469-\n1527") are preserved.
_LINE_BREAK_HYPHEN = re.compile(
    r"(?<=[^\W\d_])[ \t]*-[ \t]*\n[ \t]*(?=[^\W\d_])"
)
# Same, but a page number sits between the two halves of the word:
# "ine -\n6\n\nsauribile" -> "inesauribile".
_PAGE_BREAK_HYPHEN = re.compile(
    r"(?<=[^\W\d_])[ \t]*-[ \t]*\n\d{1,4}\n{1,2}[ \t]*(?=[^\W\d_])"
)


def normalize_pdf_text(text: str) -> str:
    """Normalize common line-break artifacts created by PDF extraction.

    Older PDFs split words across lines with a hyphen. Re-joining them keeps
    the text searchable and lets exact-quote validation match a correct
    model citation.
    """
    text = _PAGE_BREAK_HYPHEN.sub("", text)
    text = _LINE_BREAK_HYPHEN.sub("", text)
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


# Bump when text extraction or chunking changes, so stale embedding caches
# are never reused with different chunks.
INGESTION_VERSION = 3

# Every page of the bundled PDF ends with its page number on its own line.
# In the middle of a sentence that number lands inside the text, so a quote
# copied across a page break can never match the source exactly.
_TRAILING_PAGE_NUMBER = re.compile(r"\n[ \t]*\d{1,4}[ \t]*\Z")


def join_pdf_pages(pages: list[str]) -> str:
    """Join extracted PDF pages, dropping the trailing page number of each.

    A page that stops mid-sentence is joined with a single newline so that
    sentences (and hyphenated words) continue across the page break; a page
    that ends a sentence keeps a paragraph break.
    """
    result = ""

    for page in pages:
        text = _TRAILING_PAGE_NUMBER.sub("", page.rstrip())

        if not text.strip():
            continue

        if result:
            ended = re.search(r"[.!?:\u00bb\"\u201d]\s*\Z", result)
            result += "\n\n" if ended else "\n"

        result += text.rstrip()

    return result


def read_document(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md"}:
        for encoding in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                pass

        return path.read_text(encoding="utf-8", errors="replace")

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Install pypdf for PDF support.") from exc

        extracted_text = join_pdf_pages(
            [
                page.extract_text() or ""
                for page in PdfReader(str(path)).pages
            ]
        )

        return normalize_pdf_text(extracted_text)

    if suffix == ".docx":
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise RuntimeError(
                "Install python-docx for DOCX support."
            ) from exc

        document = DocxDocument(str(path))
        parts = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                )

                if row_text:
                    parts.append(row_text)

        return "\n".join(parts)

    raise ValueError(f"Unsupported file type: {suffix}")


def load_documents(root: str, debug: bool) -> list[Document]:
    base = Path(root).expanduser().resolve()

    if not base.exists():
        raise FileNotFoundError(f"Documents path not found: {base}")

    paths = (
        [base]
        if base.is_file()
        else sorted(
            path
            for path in base.rglob("*")
            if path.is_file()
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
        )
    )

    documents: list[Document] = []

    for path in paths:
        try:
            text = read_document(path).strip()

            if not text:
                log(f"[LOAD] Empty file skipped: {path.name}", debug)
                continue

            documents.append(
                Document(
                    doc_id=stable_id(str(path.resolve())),
                    path=str(path.resolve()),
                    title=path.name,
                    text=text,
                    fingerprint=stable_id(text, 40),
                )
            )

            log(
                f"[LOAD] {path.name}: {len(text):,} characters",
                debug,
            )

        except Exception as exc:
            log(f"[LOAD][ERROR] {path.name}: {exc}", debug, logging.WARNING)

    if not documents:
        raise RuntimeError("No readable documents found.")

    return documents


def make_chunks(
        documents: list[Document],
        chunk_size: int,
        overlap: int,
) -> list[Chunk]:
    if chunk_size < 100:
        raise ValueError("chunk_size must be at least 100.")

    if overlap < 0 or overlap >= chunk_size:
        raise ValueError(
            "chunk_overlap must be >= 0 and smaller than chunk_size."
        )

    chunks: list[Chunk] = []

    for document in documents:
        start = 0
        ordinal = 0
        text = document.text

        while start < len(text):
            end = min(start + chunk_size, len(text))

            if end < len(text):
                floor = start + int(chunk_size * 0.62)

                boundary = max(
                    text.rfind("\n\n", floor, end),
                    text.rfind("\n", floor, end),
                    text.rfind(". ", floor, end),
                    text.rfind("! ", floor, end),
                    text.rfind("? ", floor, end),
                )

                if boundary > floor:
                    marker = text[boundary:boundary + 2]
                    end = boundary + (
                        2
                        if marker in {"\n\n", ". ", "! ", "? "}
                        else 1
                    )

            # Never end a chunk in the middle of a word.
            if (
                end < len(text)
                and text[end - 1].isalnum()
                and text[end].isalnum()
            ):
                space = max(
                    text.rfind(" ", start, end),
                    text.rfind("\n", start, end),
                )

                if space > start + chunk_size // 2:
                    end = space + 1

            piece = text[start:end]

            if piece.strip():
                chunks.append(
                    Chunk(
                        chunk_id=f"{document.doc_id}:{ordinal}",
                        doc_id=document.doc_id,
                        text=piece,
                        start=start,
                        end=end,
                        ordinal=ordinal,
                    )
                )
                ordinal += 1

            if end >= len(text):
                break

            next_start = max(start + 1, end - overlap)

            # Start the next chunk at the beginning of a word: finish the word
            # we landed in, then skip the spaces/punctuation after it.
            while (
                next_start < end
                and text[next_start - 1].isalnum()
                and text[next_start].isalnum()
            ):
                next_start += 1

            while next_start < end and not text[next_start].isalnum():
                next_start += 1

            start = next_start

    return chunks
